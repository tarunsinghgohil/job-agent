"""Resume text extraction, section detection, and skill extraction.

Deterministic on purpose. AI is not required to read a resume, so uploads keep
working (and stay free) when no provider key is configured.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.services.normalize import norm_text

logger = logging.getLogger(__name__)

# Canonical skill vocabulary used for deterministic extraction. Matching is
# alias-aware so "ReactJS", "React.js" and "React" all resolve to one skill.
SKILL_VOCABULARY: dict[str, tuple[str, ...]] = {
    "React": ("react", "react.js", "reactjs", "react js"),
    "React Native": ("react native",),
    "Next.js": ("next.js", "nextjs", "next js"),
    "JavaScript": ("javascript", "es6", "ecmascript"),
    "TypeScript": ("typescript", "ts"),
    "HTML": ("html", "html5"),
    "CSS": ("css", "css3"),
    "Sass": ("sass", "scss"),
    "Tailwind CSS": ("tailwind", "tailwind css"),
    "Material UI": ("material ui", "mui", "material-ui"),
    "Bootstrap": ("bootstrap",),
    "Redux": ("redux",),
    "Redux Toolkit": ("redux toolkit", "rtk"),
    "Zustand": ("zustand",),
    "React Query": ("react query", "tanstack query"),
    "React Router": ("react router",),
    "GraphQL": ("graphql", "apollo"),
    "REST APIs": ("rest api", "rest apis", "restful"),
    "Node.js": ("node.js", "nodejs", "node js"),
    "Express.js": ("express.js", "express", "expressjs"),
    "Python": ("python",),
    "Django": ("django",),
    "FastAPI": ("fastapi",),
    "Flask": ("flask",),
    "Java": ("java",),
    "Go": ("golang", "go lang"),
    "PostgreSQL": ("postgresql", "postgres"),
    "MySQL": ("mysql",),
    "MongoDB": ("mongodb", "mongo"),
    "Redis": ("redis",),
    "AWS": ("aws", "amazon web services"),
    "Azure": ("azure",),
    "GCP": ("gcp", "google cloud"),
    "Docker": ("docker",),
    "Kubernetes": ("kubernetes", "k8s"),
    "CI/CD": ("ci/cd", "continuous integration", "continuous delivery"),
    "Git": ("git", "github", "gitlab", "bitbucket"),
    "Jest": ("jest",),
    "Vitest": ("vitest",),
    "React Testing Library": ("react testing library", "testing library"),
    "Cypress": ("cypress",),
    "Playwright": ("playwright",),
    "Webpack": ("webpack",),
    "Vite": ("vite",),
    "Babel": ("babel",),
    "Figma": ("figma",),
    "Accessibility": ("accessibility", "a11y", "wcag"),
    "Performance Optimization": ("performance optimization", "web vitals", "lighthouse"),
    "Agile": ("agile", "scrum", "kanban"),
    "Microservices": ("microservices",),
    "WebSockets": ("websocket", "websockets", "socket.io"),
    "Jenkins": ("jenkins",),
    "Terraform": ("terraform",),
    "OpenAI": ("openai", "gpt", "chatgpt"),
    "AI/LLM": (
        "llm", "llms", "large language model", "large language models",
        "genai", "gen ai", "generative ai", "ai/llm",
    ),
    "RAG": ("rag", "retrieval augmented generation", "retrieval-augmented generation"),
    "Angular": ("angular", "angularjs", "angular.js"),
    "Vue.js": ("vue", "vue.js", "vuejs"),
    "Storybook": ("storybook",),
    "Recharts": ("recharts",),
    "D3.js": ("d3", "d3.js"),
    "Micro Frontends": ("micro frontend", "micro frontends", "micro-frontend", "module federation"),
}

# Stack shorthands that stand for several canonical skills at once. A resume
# or posting that says "MERN" is claiming (or asking for) all four.
SKILL_GROUPS: dict[str, tuple[str, ...]] = {
    "mern": ("MongoDB", "Express.js", "React", "Node.js"),
    "mern stack": ("MongoDB", "Express.js", "React", "Node.js"),
    "mean": ("MongoDB", "Express.js", "Angular", "Node.js"),
    "mean stack": ("MongoDB", "Express.js", "Angular", "Node.js"),
    "mevn": ("MongoDB", "Express.js", "Vue.js", "Node.js"),
    "pern": ("PostgreSQL", "Express.js", "React", "Node.js"),
}

# Extra aliases that are only safe when the whole input *is* the alias, e.g. a
# skill chip typed as "JS". Matching "js" inside running text would fire on
# every "Node.js" and "Next.js".
_EXACT_ONLY_ALIASES: dict[str, str] = {
    "js": "JavaScript",
    "ecma": "JavaScript",
    "node": "Node.js",
    "next": "Next.js",
    "express": "Express.js",
    "postgres/mysql": "PostgreSQL",
    "postgresql/mysql": "PostgreSQL",
}

# Headings that mark a section. Order does not matter; position in the text does.
SECTION_PATTERNS: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "profile", "objective", "about me", "professional summary"),
    "experience": (
        "experience", "work experience", "employment", "professional experience",
        "work history", "career history",
    ),
    "education": ("education", "academic", "qualifications", "academics"),
    "skills": ("skills", "technical skills", "core competencies", "technologies", "tech stack"),
    "projects": ("projects", "personal projects", "key projects", "selected projects"),
    "certifications": ("certifications", "certificates", "licenses", "courses"),
    "achievements": ("achievements", "awards", "honors", "accomplishments"),
    "contact": ("contact", "contact information", "personal details"),
}

_MAX_HEADING_WORDS = 6


@dataclass(slots=True)
class ParsedResume:
    text: str = ""
    sections: dict[str, str] = field(default_factory=dict)
    skills: list[str] = field(default_factory=list)
    emails: list[str] = field(default_factory=list)
    phones: list[str] = field(default_factory=list)
    links: list[str] = field(default_factory=list)
    error: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.text.strip()) and not self.error


def extract_text(path: str | Path, extension: str = "") -> str:
    """Extract plain text from a PDF or DOCX file."""
    path = Path(path)
    ext = (extension or path.suffix).lower()

    if ext == ".pdf":
        import fitz  # PyMuPDF

        parts: list[str] = []
        with fitz.open(path) as doc:
            for page in doc:
                parts.append(page.get_text("text"))
        return "\n".join(parts)

    if ext == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        # Table cells often hold skills grids, which would otherwise be lost.
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    raise ValueError(f"Unsupported resume format: {ext!r}")


def _looks_like_heading(line: str) -> str | None:
    """Return the canonical section name when a line reads as its heading."""
    stripped = line.strip().strip(":").strip()
    if not stripped or len(stripped.split()) > _MAX_HEADING_WORDS:
        return None
    normalized = norm_text(stripped)
    if not normalized:
        return None
    for section, aliases in SECTION_PATTERNS.items():
        if normalized in aliases:
            return section
    return None


def detect_sections(text: str) -> dict[str, str]:
    """Split resume text into named sections by heading lines."""
    lines = text.splitlines()
    sections: dict[str, list[str]] = {}
    current = "header"

    for line in lines:
        heading = _looks_like_heading(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)

    return {
        name: "\n".join(body).strip()
        for name, body in sections.items()
        if "\n".join(body).strip()
    }


_ALIAS_PATTERNS: dict[str, re.Pattern[str]] = {}


def _alias_pattern(needle: str) -> re.Pattern[str]:
    """Whole-word pattern for a normalized alias.

    Plain substring matching made short aliases fire inside ordinary words:
    "ts" matched "requirements", "git" matched "digital", "java" matched
    "javascript". Word boundaries are expressed as "not alphanumeric" so
    aliases that contain dots or pluses ("node.js", "c++") still work.
    """
    pattern = _ALIAS_PATTERNS.get(needle)
    if pattern is None:
        pattern = re.compile(rf"(?<![a-z0-9]){re.escape(needle)}(?![a-z0-9])")
        _ALIAS_PATTERNS[needle] = pattern
    return pattern


def contains_term(term: str, text: str, *, normalized: bool = False) -> bool:
    """True when ``term`` appears in ``text`` as a whole word or phrase."""
    needle = norm_text(term)
    if not needle:
        return False
    haystack = text if normalized else norm_text(text)
    return bool(_alias_pattern(needle).search(haystack))


def extract_skills(text: str) -> list[str]:
    """Find canonical skills present in the text, preserving vocabulary order.

    Stack shorthands such as "MERN" expand into their member skills.
    """
    haystack = norm_text(text)
    found: list[str] = []
    for canonical, aliases in SKILL_VOCABULARY.items():
        for alias in aliases:
            needle = norm_text(alias)
            if needle and _alias_pattern(needle).search(haystack):
                found.append(canonical)
                break
    for shorthand, members in SKILL_GROUPS.items():
        if _alias_pattern(shorthand).search(haystack):
            for member in members:
                if member not in found:
                    found.append(member)
    return found


def canonicalize_skill(name: str) -> list[str]:
    """Map one user-typed skill to canonical skill names.

    "ReactJS" -> ["React"], "MERN" -> the four MERN skills, "JS" ->
    ["JavaScript"]. A skill outside the vocabulary is kept as typed so custom
    skills ("Recoil", "HIMS") still participate in matching by name.
    """
    raw = (name or "").strip()
    key = norm_text(raw)
    if not key:
        return []
    if key in SKILL_GROUPS:
        return list(SKILL_GROUPS[key])
    if key in _EXACT_ONLY_ALIASES:
        return [_EXACT_ONLY_ALIASES[key]]
    for canonical, aliases in SKILL_VOCABULARY.items():
        if key == norm_text(canonical) or any(key == norm_text(a) for a in aliases):
            return [canonical]
    # Composite chips like "PostgreSQL/MySQL" or "React + Redux".
    parts = [p for p in re.split(r"\s*(?:/|\+|,|&|\band\b)\s*", raw) if p.strip()]
    if len(parts) > 1:
        out: list[str] = []
        for part in parts:
            for item in canonicalize_skill(part):
                if item not in out:
                    out.append(item)
        return out
    return [raw]


def normalize_skill_list(names: list[str] | None) -> list[str]:
    """Canonicalize and de-duplicate a list of skills, preserving order."""
    out: list[str] = []
    seen: set[str] = set()
    for name in names or []:
        for canonical in canonicalize_skill(str(name)):
            key = canonical.lower()
            if key not in seen:
                seen.add(key)
                out.append(canonical)
    return out


_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?:\+\d{1,3}[\s-]?)?(?:\d[\s-]?){9,13}\d")
_URL_RE = re.compile(r"(?:https?://|www\.)[^\s,;)\]]+", re.I)


def parse(path: str | Path, extension: str = "") -> ParsedResume:
    """Full parse. Never raises: failures come back on ``error``."""
    try:
        text = extract_text(path, extension)
    except Exception as exc:
        logger.warning("Resume text extraction failed for %s: %s", path, exc)
        return ParsedResume(error=f"Could not read the document: {exc}")

    if not text.strip():
        return ParsedResume(
            text="",
            error=(
                "No text could be extracted. The file is most likely a scanned "
                "image; a text-based PDF or DOCX is required."
            ),
        )

    return ParsedResume(
        text=text,
        sections=detect_sections(text),
        skills=extract_skills(text),
        emails=sorted(set(_EMAIL_RE.findall(text)))[:5],
        phones=sorted({re.sub(r"[\s-]", "", p) for p in _PHONE_RE.findall(text) if len(re.sub(r"[^0-9]", "", p)) >= 10})[:3],
        links=sorted(set(_URL_RE.findall(text)))[:10],
    )
